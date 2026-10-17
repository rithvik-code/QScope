from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "QScope_Project_Document.docx"
NAVY = "17324D"
TEAL = "0F766E"
INK = "1F2937"
MUTED = "5B6472"
PALE = "EAF4F3"
LINE = "D7DEE7"


def shade(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), colour)


def set_cell_border(cell, colour=LINE, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), colour)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def section_heading(doc, number, title):
    p = doc.add_paragraph(style="Section Heading")
    p.add_run(f"{number}. ").font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(title)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        shade(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2 == 0:
                shade(cells[idx], "F7FAFC")
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor.from_string(INK)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size, colour in (("Title", 29, NAVY), ("Subtitle", 12, TEAL)):
    styles[name].font.name = "Aptos Display"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(colour)
styles["Title"].font.bold = True
styles["Title"].paragraph_format.space_after = Pt(4)
styles["Subtitle"].paragraph_format.space_after = Pt(14)
if "Section Heading" not in styles:
    styles.add_style("Section Heading", WD_STYLE_TYPE.PARAGRAPH)
styles["Section Heading"].base_style = styles["Heading 1"]
styles["Section Heading"].font.name = "Aptos Display"
styles["Section Heading"].font.size = Pt(16)
styles["Section Heading"].font.bold = True
styles["Section Heading"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Section Heading"].paragraph_format.space_before = Pt(12)
styles["Section Heading"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.size = Pt(11.5)
styles["Heading 2"].font.color.rgb = RGBColor.from_string(TEAL)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run("QSCOPE  /  QUANTUM COMPUTING WORKSPACE")
header_run.font.size = Pt(8)
header_run.font.bold = True
header_run.font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("QScope · Local-first quantum computing workspace   |   ")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string(MUTED)
add_field(footer, "PAGE")

title = doc.add_paragraph(style="Title")
title.add_run("QScope")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("Local-first quantum computing workspace")

student = add_table(doc, ["Name", "Register number"], [["K Rithvik", "RA2511003010883"]], [3.45, 3.5])
for cell in student.rows[1].cells:
    shade(cell, PALE)

meta = add_table(doc, ["Project", "Category", "Operating model", "Version"], [[
    "QScope — See inside quantum computing",
    "Quantum-computing education and research workspace",
    "Local-first web application",
    "0.1.0",
]], [2.2, 2.25, 1.6, 0.8])
for cell in meta.rows[1].cells:
    shade(cell, PALE)

callout = doc.add_table(rows=1, cols=1)
callout.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = callout.cell(0, 0)
shade(cell, PALE)
set_cell_border(cell, "9CCFC8", "10")
set_cell_margins(cell, 150, 180, 150, 180)
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Ask a question. Build or choose a circuit. Run it locally. See what happened. Understand why.")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(NAVY)
doc.add_paragraph().paragraph_format.space_after = Pt(1)

section_heading(doc, "1", "Abstract")
abstract = ("QScope is a local-first quantum-computing workspace that helps learners and researchers move from a plain-language question to a runnable, explainable experiment. It combines a circuit builder, OpenQASM import, reusable algorithm templates, classical simulation engines, gate-by-gate tracing, state analysis, optimization, noise modelling, reproducible experiments, benchmarking, history, reporting, and a grounded assistant in one browser-based application. The system runs on an ordinary laptop using a Python backend and a React-based frontend, with no cloud quantum account required for its core workflow. QScope emphasizes transparency: probabilities, amplitudes, entanglement, fidelity, memory, runtime, seeds, and execution modes remain attached to the results that produced them. Its guided Build–Run–Understand loop makes quantum concepts approachable while progressive disclosure preserves advanced research controls. By connecting education, experimentation, visualization, and provenance, QScope provides a practical environment for understanding what a quantum circuit does, why its output changes, and how to reproduce or extend the investigation responsibly in real settings.")
assert len(abstract.split()) == 150, len(abstract.split())
doc.add_paragraph(abstract)

section_heading(doc, "2", "Problem Definition")
doc.add_paragraph("Quantum computing is difficult to learn and investigate because the user must often combine several disconnected tools: a circuit editor, a simulator, a visualization package, a debugging method, an experiment log, and explanatory documentation. Traditional interfaces frequently expose advanced controls before establishing a clear starting point, while simplified demonstrations hide the intermediate states and assumptions that make results meaningful.")
doc.add_paragraph("The central problem addressed by QScope is therefore: how can a learner or researcher construct, execute, inspect, explain, compare, and reproduce quantum-circuit experiments in one approachable environment, while preserving technical accuracy and clearly distinguishing classical simulation from real quantum hardware?")

section_heading(doc, "3", "Objectives")
for item in [
    "Provide a simple entry point that lets users begin with a goal, a template, imported QASM, or a blank circuit.",
    "Enable local execution of ideal and noisy quantum circuits without requiring a cloud account, quantum processor, GPU, or paid service.",
    "Make intermediate behavior visible through traces, probability distributions, amplitudes, phase information, entanglement, and fidelity metrics.",
    "Support reproducible work through explicit settings, seeds, stored runs, circuit snapshots, comparison, and report export.",
    "Offer a grounded assistant that explains the user’s actual circuits and results, identifies missing evidence, and avoids inventing numerical claims.",
    "Expose advanced analysis, optimization, experimentation, benchmarking, hardware estimates, and research features through progressive disclosure.",
]: bullet(doc, item)

section_heading(doc, "4", "Scope")
add_table(doc, ["Included in scope", "Boundary / limitation"], [
    ("Circuit construction, validation, editing, templates, algorithms, and OpenQASM 2.0 import/export.", "The core engine is a classical simulator; it does not claim quantum advantage."),
    ("State-vector, density-matrix, and trajectory simulation with configurable shots, seeds, and noise.", "Circuit size is limited by the laptop’s available memory and compute budget."),
    ("Visual debugging, analysis, optimization, experiments, benchmarks, history, reports, and assistant explanations.", "Hardware values are estimates unless an explicitly connected provider is used."),
    ("Local web interface served at 127.0.0.1:8000 with optional local development server support.", "Optional LLM narration depends on a configured OpenAI-compatible endpoint."),
], [3.35, 3.5])

section_heading(doc, "5", "Proposed System")
doc.add_paragraph("QScope proposes a unified, local-first quantum workspace organized around the continuous loop:")
add_table(doc, ["Stage", "User action", "System response"], [
    ("Choose a starting point", "Select a learning goal, template, algorithm, QASM file, or blank circuit.", "Loads an explicit circuit document with readable metadata and a visual preview."),
    ("Build", "Add gates, edit qubits, configure sources, and inspect circuit resources.", "Validates the circuit and keeps the interface synchronized with the Python engine."),
    ("Run", "Choose shots, seed, backend, noise, and history settings.", "Produces a labelled simulation result with provenance, timing, memory, counts, and probabilities."),
    ("Understand", "Trace gates, inspect state changes, ask questions, analyze, optimize, or compare.", "Shows intermediate states and plain-language explanations grounded in recorded artifacts."),
    ("Return to it", "Restore, reproduce, compare, export, or extend a previous run.", "Uses stored circuit snapshots, settings, results, and experiment history."),
], [1.35, 2.5, 3.0])

section_heading(doc, "6", "Existing Approaches")
doc.add_paragraph("Existing quantum-computing tools generally fall into separate categories: educational circuit visualizers, programming libraries, cloud quantum platforms, notebook-based simulators, and research-focused benchmarking or analysis tools. Each category is useful, but users commonly need to move data manually between them. Beginner tools may provide attractive output without reproducibility or technical depth; research tools may expose depth without a calm onboarding path; cloud platforms may introduce accounts, network dependence, and hardware terminology before the user understands the underlying circuit.")
doc.add_paragraph("QScope does not attempt to replace specialized quantum hardware providers or mature scientific libraries. Instead, it integrates the essential learning and research loop in a single local environment and keeps its boundaries explicit.")

section_heading(doc, "7", "How QScope Is Better")
add_table(doc, ["Design advantage", "Practical benefit"], [
    ("One continuous workspace", "The same circuit and recorded context flow from Build to Run, Trace, Analyze, Optimize, Experiment, Benchmark, and Research."),
    ("Local-first operation", "Users can learn and prototype immediately on a normal laptop without an account or quantum cloud provider."),
    ("Explainability by default", "Results include intermediate state changes and plain-language interpretation instead of only final counts."),
    ("Evidence-grounded assistance", "The assistant identifies its sources and refuses to invent missing numerical evidence."),
    ("Progressive disclosure", "A calm beginner path coexists with advanced controls for noise, optimization, experiments, and performance."),
    ("Reproducibility and provenance", "Seeds, shots, backend, noise, timing, memory, settings, circuit snapshots, and stored history travel with the result."),
], [2.2, 4.65])

section_heading(doc, "8", "Technical Coverage")
doc.add_paragraph("The system is implemented as a local web application with a Python source-of-truth backend and a React/TypeScript interface. The following coverage describes the current project boundary.")

doc.add_paragraph("Backend", style="Heading 2")
for item in [
    "Python 3.11+ application organized into circuit, core, simulator, debugger, analysis, optimizer, algorithms, experiments, hardware, evolution, AI, and API packages.",
    "FastAPI service served by Uvicorn, with REST endpoints for health, metadata, circuits, simulation, traces, analysis, optimization, experiments, history, reports, hardware estimates, and assistant requests.",
    "NumPy-based state-vector, density-matrix, and trajectory simulation engines with explicit ideal/noisy execution modes.",
    "Pydantic models and validation for circuit documents, run settings, API requests, results, provenance, and stored artifacts.",
    "WebSocket support for streamed traces and experiment progress where interactive feedback is useful.",
    "SQLite-backed local experiment history with stored circuit snapshots, run settings, results, reproduction, comparison, and CSV/JSON export.",
    "Grounded AI assistant with intent classification, evidence bundles, citations, limitations, suggestions, and optional OpenAI-compatible narration.",
]: bullet(doc, item)

doc.add_paragraph("Frontend", style="Heading 2")
for item in [
    "React and TypeScript single-page interface bundled with Vite.",
    "Tailwind-style utility classes and project-specific CSS variables for responsive layout, themes, panels, tags, buttons, notices, and accessibility focus states.",
    "Mission Control onboarding plus progressive Beginner and Research-detail views.",
    "Circuit IDE with gate palette, source selection, guided templates, QASM import, circuit canvas, selection state, and resource summaries.",
    "Execute view with run settings, execution planning, result summaries, measurement counts, probability views, noisy-versus-ideal comparison, and explain actions.",
    "Trace, Analyze, Optimize, Noise, Experiment, Benchmark, History, and Research modes connected through shared application state.",
    "Floating Ask QScope local chat with session conversation history, suggested prompts, circuit/run context, loading feedback, and grounded responses.",
]: bullet(doc, item)

doc.add_paragraph("Languages, APIs, and supporting technologies", style="Heading 2")
add_table(doc, ["Area", "Technology / interface", "Role"], [
    ("Backend language", "Python 3.11+", "Simulation, domain logic, validation, persistence, reporting, and API service."),
    ("Frontend language", "TypeScript + JSX", "Typed React components, application state, interaction, and visualization."),
    ("Numerical computing", "NumPy", "State representations, matrix operations, probabilities, metrics, and simulation."),
    ("Web API", "FastAPI + Uvicorn", "Local HTTP API, OpenAPI documentation, request validation, and server runtime."),
    ("Streaming", "WebSocket", "Progressive traces and long-running experiment updates."),
    ("Build tooling", "Vite + npm", "Frontend development server, type checking, and production bundle."),
    ("Persistence", "SQLite", "Local experiment history, stored runs, circuit snapshots, and reproducibility metadata."),
    ("Data formats", "JSON, CSV, OpenQASM 2.0", "API payloads, exports, and quantum-circuit interoperability."),
    ("Optional narration", "OpenAI-compatible /chat/completions endpoint", "Natural-language rewriting of already grounded answers; never the source of numerical truth."),
    ("Testing", "pytest, HTTPX, TypeScript checks, Vite build", "Backend verification, API checks, frontend validation, and production-build confidence."),
], [1.5, 2.35, 3.3])

section_heading(doc, "9", "Expected Outcome")
doc.add_paragraph("The completed QScope system will provide a credible, demonstrable project that combines quantum-computing education with practical software engineering. A first-time user can choose a starting point, build or load a circuit, run a local simulation, see the state change gate by gate, ask a plain-language question, and return to the saved work. An advanced user can inspect metrics, noise, optimization, experiments, benchmarks, history, and exports without leaving the workspace. This combination makes the project useful as both an approachable learning tool and a foundation for reproducible quantum-circuit investigation.")

section_heading(doc, "10", "Conclusion")
doc.add_paragraph("QScope addresses the gap between approachable quantum demonstrations and technically responsible experimentation. Its defining contribution is not only the simulator, but the connected workflow around it: guided entry, visible computation, grounded explanation, and reproducible return. The architecture keeps the Python engine authoritative, the frontend focused on interaction and presentation, and optional AI narration subordinate to recorded evidence. This gives the project a clear identity, a realistic implementation boundary, and a strong basis for continued development.")

doc.save(OUT)
print(OUT)
