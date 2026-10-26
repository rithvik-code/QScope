"""Add five project-specific figures to the user's edited QScope manuscript."""

from __future__ import annotations

import io
import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "QScope_Research_Paper_Rithvik.docx"
OUTPUT = ROOT / "QScope_Research_Paper_Rithvik_Illustrated_v7.docx"
FIGURE_DIR = ROOT / "docs" / "research_figures"
W, H = 1800, 1120

INK = "#16243A"
MUTED = "#60738B"
TEAL = "#12B89C"
BLUE = "#4E8EF7"
VIOLET = "#8069E8"
AMBER = "#E4A13A"
PALE = "#F2F6FA"
LINE = "#D9E2EC"
WHITE = "#FFFFFF"
DARK = "#0C111B"
DARK_PANEL = "#141C29"
DARK_LINE = "#2A3547"


def font(size: int, weight: str = "regular") -> ImageFont.FreeTypeFont:
    names = {
        "regular": "segoeui.ttf",
        "semibold": "seguisb.ttf",
        "bold": "segoeuib.ttf",
        "mono": "consola.ttf",
    }
    path = Path("C:/Windows/Fonts") / names[weight]
    return ImageFont.truetype(str(path), size=size)


def canvas(*, dark: bool = False, label: str = "Q S C O P E   /   RESEARCH FIGURE"):
    im = Image.new("RGB", (W, H), DARK if dark else WHITE)
    d = ImageDraw.Draw(im)
    fg = "#F3F6FA" if dark else INK
    sub = "#9AA9BA" if dark else MUTED
    d.text((76, 54), label, fill=TEAL, font=font(22, "semibold"))
    return im, d, fg, sub


def rounded(draw, box, fill, outline=None, radius=24, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def wrap(draw, text, fnt, max_width):
    words, lines, line = text.split(), [], ""
    for word in words:
        candidate = (line + " " + word).strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def paragraph(draw, xy, text, fnt, fill, max_width, line_gap=8):
    x, y = xy
    for line in wrap(draw, text, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def title(draw, text, *, dark=False):
    draw.text((76, 100), text, fill="#F3F6FA" if dark else INK, font=font(46, "bold"))


def save_image(image: Image.Image, name: str) -> bytes:
    output = io.BytesIO()
    image.save(output, format="PNG", dpi=(240, 240), optimize=True)
    data = output.getvalue()
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    (FIGURE_DIR / name).write_bytes(data)
    return data


def fig_workspace() -> bytes:
    im, d, fg, sub = canvas(dark=True, label="Q S C O P E   /   LOCAL WORKSPACE OVERVIEW")
    title(d, "A quantum workspace, built around understanding", dark=True)
    rounded(d, (62, 194, 1738, 1060), "#0A0F17", "#263345", 28, 3)
    d.rectangle((64, 196, 1736, 286), fill="#111925")
    d.ellipse((98, 228, 128, 258), outline="#55D8BA", width=4)
    d.ellipse((108, 238, 118, 248), fill="#55D8BA")
    d.text((142, 220), "QScope", font=font(34, "bold"), fill="#F2F5F9")
    d.text((98, 258), "SEE INSIDE QUANTUM COMPUTING", font=font(16, "semibold"), fill="#93A4B7")
    for x, label in [(465, "MISSION CONTROL"), (680, "BUILD"), (825, "EXECUTE"), (1000, "TRACE"), (1140, "HISTORY"), (1305, "RESEARCH")]:
        d.text((x, 238), label, font=font(17, "semibold"), fill="#C6D0DB" if label == "MISSION CONTROL" else "#8090A3")
    rounded(d, (1515, 218, 1684, 264), "#143D38", "#1D7564", 18, 2)
    d.text((1540, 230), "OFFLINE · LOCAL", font=font(15, "semibold"), fill="#62D9BB")

    d.text((100, 330), "Local quantum workspace", font=font(18, "semibold"), fill="#52D9BA")
    d.text((100, 367), "What do you want to understand?", font=font(39, "bold"), fill=fg)
    d.text((102, 425), "Ask a question, choose an experiment, or start with a blank circuit.", font=font(21), fill="#AAB7C7")
    rounded(d, (100, 474, 1665, 540), "#111925", "#344256", 18, 2)
    d.text((126, 493), "Try: show me how entanglement works", font=font(19), fill="#728196")
    rounded(d, (1460, 483, 1645, 531), "#13B99B", None, 14)
    d.text((1502, 495), "Explore  →", font=font(20, "semibold"), fill="#06221D")

    # Status card reflects the live page: a Bell circuit is loaded, not yet run.
    rounded(d, (100, 574, 1665, 680), "#151E2B", "#2D3A4C", 22, 2)
    d.text((130, 594), "YOUR WORKSPACE", font=font(16, "semibold"), fill="#8FA1B6")
    for x, label, value in [(130, "ENGINE", "ready"), (510, "MODE", "offline"), (890, "CIRCUIT", "2 qubits"), (1270, "LAST RUN", "not yet")]:
        d.text((x, 630), label, font=font(15, "semibold"), fill="#738399")
        d.text((x, 651), value, font=font(20, "semibold"), fill="#EAF0F6")

    rounded(d, (100, 710, 920, 1015), "#141D2A", "#2D3A4C", 22, 2)
    d.text((132, 738), "CIRCUIT WORKSPACE", font=font(17, "semibold"), fill="#A7B4C4")
    d.text((132, 770), "bell_phi_plus", font=font(20, "mono"), fill="#F2F5F9")
    # Two circuit wires and the exact H -> CNOT Bell construction shown by QScope.
    for y, q in [(852, "q0"), (920, "q1")]:
        d.text((145, y - 15), q, font=font(18, "mono"), fill="#90A1B5")
        d.line((215, y, 865, y), fill="#435166", width=3)
    rounded(d, (350, 826, 435, 878), "#203B3C", "#36C8AC", 10, 2)
    d.text((380, 835), "H", font=font(22, "bold"), fill="#77E7D0")
    rounded(d, (570, 826, 686, 945), "#24273D", "#8C7AFF", 12, 2)
    d.text((590, 865), "CNOT", font=font(19, "bold"), fill="#C7BDFF")
    d.ellipse((621, 842, 638, 859), fill="#B9AEFF")
    d.ellipse((621, 910, 638, 927), outline="#B9AEFF", width=4)
    d.line((629, 859, 629, 910), fill="#B9AEFF", width=3)
    for yy in (827, 897):
        rounded(d, (765, yy, 820, yy + 50), "#1C2D3C", "#5A7894", 9, 2)
        d.text((780, yy + 12), "M", font=font(20, "bold"), fill="#BBD0E4")
    d.text((132, 976), "H on q0  →  CNOT(q0, q1)  →  measure", font=font(17), fill="#93A2B5")

    rounded(d, (950, 710, 1665, 1015), "#141D2A", "#2D3A4C", 22, 2)
    d.text((985, 738), "QUANTUM STATE", font=font(17, "semibold"), fill="#A7B4C4")
    d.text((985, 795), "Run a circuit to see its state.", font=font(25, "semibold"), fill="#EDF2F8")
    d.text((985, 841), "Probability map and phase view appear after a run.", font=font(18), fill="#95A4B7")
    rounded(d, (985, 897, 1628, 965), "#101722", "#253447", 15, 2)
    d.text((1010, 918), "BUILD", font=font(16, "semibold"), fill="#57D4B7")
    d.text((1165, 918), "RUN", font=font(16, "semibold"), fill="#79A9FF")
    d.text((1310, 918), "UNDERSTAND", font=font(16, "semibold"), fill="#B2A5FF")
    d.text((1010, 946), "Create/import", font=font(14), fill="#8596AA")
    d.text((1165, 946), "Simulate locally", font=font(14), fill="#8596AA")
    d.text((1310, 946), "Trace every gate", font=font(14), fill="#8596AA")
    return save_image(im, "qscope-figure-01-workspace-v3.png")


def fig_architecture() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   SYSTEM ARCHITECTURE")
    title(d, "One local workflow, four cooperating layers")
    rounded(d, (75, 220, 1725, 990), "#F7FAFD", "#CFDCE8", 30, 3)
    d.text((110, 250), "USER'S LAPTOP  ·  LOCAL-FIRST  ·  NO QUANTUM HARDWARE REQUIRED", font=font(20, "semibold"), fill=TEAL)
    cards = [
        (112, 366, 478, 714, "01  INTERFACE", "React + TypeScript", ["Circuit editor", "Guided experiments", "State and trace views", "History and reports"], BLUE),
        (536, 366, 902, 714, "02  LOCAL API", "FastAPI + Pydantic", ["Validate circuits", "Plan execution", "Serve run/trace APIs", "Normalize QASM input"], TEAL),
        (960, 366, 1326, 714, "03  QUANTUM ENGINE", "Python + NumPy", ["State-vector", "Density matrix", "Noisy trajectories", "Analysis + optimization"], VIOLET),
        (1384, 366, 1689, 714, "04  PERSISTENCE", "SQLite + local files", ["Run history", "Versioned drafts", "Project bundles", "Reproducibility data"], AMBER),
    ]
    for x1, y1, x2, y2, head, subhead, bullets, accent in cards:
        rounded(d, (x1, y1, x2, y2), WHITE, LINE, 24, 2)
        d.rounded_rectangle((x1, y1, x2, y1 + 15), radius=8, fill=accent)
        d.text((x1 + 24, y1 + 42), head, font=font(18, "bold"), fill=INK)
        d.text((x1 + 24, y1 + 82), subhead, font=font(17, "semibold"), fill=accent)
        yy = y1 + 142
        for bullet in bullets:
            d.ellipse((x1 + 26, yy + 7, x1 + 38, yy + 19), fill=accent)
            d.text((x1 + 52, yy), bullet, font=font(17), fill="#42566D")
            yy += 49
    # Request flow and shared data persistence.
    for x1, x2 in [(480, 530), (904, 954)]:
        d.line((x1, 520, x2, 520), fill="#7A8EA4", width=5)
        d.polygon([(x2, 520), (x2 - 16, 508), (x2 - 16, 532)], fill="#7A8EA4")
    d.line((1328, 520, 1378, 520), fill="#7A8EA4", width=5)
    d.polygon([(1378, 520), (1362, 508), (1362, 532)], fill="#7A8EA4")
    rounded(d, (112, 786, 1689, 920), "#E9F6F3", "#B8E4D9", 22, 2)
    d.text((150, 807), "INSPECTABLE EXPERIMENT LOOP", font=font(18, "bold"), fill="#168E78")
    labels = ["Build / import", "Preflight plan", "Simulate", "Trace / analyze", "Save & compare"]
    xs = [150, 470, 790, 1110, 1430]
    for i, (x, label) in enumerate(zip(xs, labels)):
        rounded(d, (x, 849, x + 205, 897), WHITE, "#B8D9D1", 15, 2)
        d.text((x + 20, 862), label, font=font(18, "semibold"), fill=INK)
        if i < len(labels) - 1:
            d.line((x + 211, 873, x + 293, 873), fill=TEAL, width=4)
            d.polygon([(x + 293, 873), (x + 280, 864), (x + 280, 882)], fill=TEAL)
    return save_image(im, "qscope-figure-02-architecture-v3.png")


def fig_modes() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   QUANTUM SIMULATION")
    title(d, "Three representations, three different trade-offs")
    cards = [
        (78, 226, 610, 966, "STATE VECTOR", "|psi>", BLUE),
        (634, 226, 1166, 966, "DENSITY MATRIX", "rho", VIOLET),
        (1190, 226, 1722, 966, "TRAJECTORIES", "{ |psi_k> }", TEAL),
    ]
    for x1, y1, x2, y2, label, symbol, accent in cards:
        rounded(d, (x1, y1, x2, y2), "#FBFCFE", "#D8E2EC", 26, 2)
        d.rounded_rectangle((x1, y1, x2, y1 + 14), radius=8, fill=accent)
        d.text((x1 + 30, y1 + 40), label, font=font(22, "bold"), fill=INK)
        d.text((x1 + 32, y1 + 95), symbol, font=font(38, "bold"), fill=accent)
    # Amplitudes
    x0, y0, x1, y1 = 125, 414, 562, 653
    d.line((x0, y1, x1, y1), fill="#AEBCCB", width=2)
    heights = [0, 10, 14, 0, 0, 0, 11, 0, 0, 0, 0, 0, 0, 0, 10, 0]
    for i, h in enumerate(heights):
        xx = x0 + 20 + i * 26
        if h:
            d.rounded_rectangle((xx, y1 - h * 11, xx + 16, y1), radius=5, fill=accent if i in (0, 15) else BLUE)
    d.text((124, 680), "Amplitudes over computational basis states", font=font(17), fill=MUTED)
    d.text((126, 740), "Good for", font=font(17, "bold"), fill=INK)
    paragraph(d, (126, 775), "Ideal pure-state evolution and sampled pure-state trajectories.", font(18), "#43566B", 420, 8)
    rounded(d, (126, 872, 562, 928), "#EAF2FF", None, 14)
    d.text((147, 889), "MEMORY   O(2^n) complex amplitudes", font=font(17, "semibold"), fill="#376BC2")
    # Density matrix heatmap
    mx, my, size, cells = 694, 407, 31, 12
    for r in range(cells):
        for c in range(cells):
            dist = abs(r - c)
            alpha = max(0, 200 - dist * 22 - abs(r + c - cells) * 2)
            color = (min(255, 226 - alpha // 3), min(255, 224 - alpha // 2), min(255, 247 - alpha // 2))
            d.rectangle((mx + c * size, my + r * size, mx + (c + 1) * size - 3, my + (r + 1) * size - 3), fill=color)
    d.text((694, 790), "Rows × columns encode coherences", font=font(17), fill=MUTED)
    d.text((694, 825), "Good for", font=font(17, "bold"), fill=INK)
    paragraph(d, (694, 855), "Exact mixed-state evolution with supported noise.", font(18), "#43566B", 420, 8)
    rounded(d, (694, 916, 1134, 958), "#F0EDFF", None, 12)
    d.text((711, 927), "MEMORY   O(4^n) complex entries", font=font(17, "semibold"), fill="#6552C6")
    # Trajectory branches, each a stochastic pure-state path
    d.text((1230, 407), "Noise channel", font=font(17, "semibold"), fill=MUTED)
    d.ellipse((1260, 458, 1318, 516), fill="#E3F6F2", outline=TEAL, width=3)
    d.text((1278, 473), "ρ", font=font(27, "bold"), fill="#168E78")
    for k, yy in enumerate([435, 515, 595]):
        d.line((1318, 487, 1402, yy + 21), fill="#89A2B6", width=3)
        d.ellipse((1400, yy, 1444, yy + 44), fill="#E9F0FF", outline=BLUE, width=2)
        d.text((1411, yy + 9), str(k + 1), font=font(19, "bold"), fill="#376BC2")
        d.line((1444, yy + 22, 1500, yy + 22), fill="#89A2B6", width=3)
        for j in range(3):
            d.ellipse((1502 + j * 45, yy + 12, 1522 + j * 45, yy + 32), fill=TEAL if (j + k) % 2 else VIOLET)
    d.text((1230, 680), "Sample branches, then aggregate shots", font=font(17), fill=MUTED)
    d.text((1230, 740), "Good for", font=font(17, "bold"), fill=INK)
    paragraph(d, (1230, 775), "Approximate noisy evolution when exact density matrices are too costly.", font(18), "#43566B", 425, 8)
    rounded(d, (1230, 872, 1678, 928), "#E7F7F3", None, 14)
    d.text((1250, 889), "COST   More shots improve sampling", font=font(16, "semibold"), fill="#168E78")
    d.text((80, 1022), "Classical simulation only · output type and approximation are identified per run.", font=font(17, "semibold"), fill=MUTED)
    return save_image(im, "qscope-figure-03-simulation-modes-v3.png")


def fig_scheduler() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   MEMORY-AWARE BATCH EXECUTION")
    title(d, "Parallelize independent circuits—within a shared budget")
    rounded(d, (90, 225, 1710, 395), "#F7FAFD", LINE, 26, 2)
    d.text((128, 253), "INDEPENDENT INPUT JOBS", font=font(18, "bold"), fill=MUTED)
    job_specs = [("A", "3 units", BLUE), ("B", "4 units", VIOLET), ("C", "5 units", TEAL), ("D", "3 units", AMBER)]
    for i, (name, mem, color) in enumerate(job_specs):
        x = 142 + i * 390
        rounded(d, (x, 301, x + 300, 364), WHITE, color, 16, 3)
        d.text((x + 24, 318), f"Job {name}", font=font(21, "bold"), fill=INK)
        d.text((x + 132, 322), f"estimated peak  {mem}", font=font(16), fill=MUTED)
    d.line((895, 395, 895, 465), fill="#8397AB", width=5)
    d.polygon([(895, 466), (880, 445), (910, 445)], fill="#8397AB")
    rounded(d, (576, 470, 1214, 570), "#EAF7F4", "#A8DED2", 24, 2)
    d.text((685, 489), "PREFLIGHT PLANNER + WAVE SCHEDULER", font=font(22, "bold"), fill="#147C6A")
    d.text((691, 528), "Worker cap + sum of planned peaks ≤ shared budget", font=font(17, "semibold"), fill="#3B766C")
    # budget guide and two waves
    d.text((115, 626), "SHARED MEMORY BUDGET  B = 8 illustrative units", font=font(19, "bold"), fill=INK)
    d.text((1274, 626), "Each wave stays within B", font=font(17, "semibold"), fill=MUTED)
    for wave, y, jobs, total in [("WAVE 1", 692, [("A", 3, BLUE), ("B", 4, VIOLET)], "3 + 4 = 7 ≤ 8"),
                                ("WAVE 2", 841, [("C", 5, TEAL), ("D", 3, AMBER)], "5 + 3 = 8 ≤ 8")]:
        rounded(d, (105, y, 1695, y + 118), WHITE, LINE, 20, 2)
        d.text((135, y + 18), wave, font=font(18, "bold"), fill=MUTED)
        xx = 310
        for name, amount, color in jobs:
            barw = 116 * amount
            rounded(d, (xx, y + 28, xx + barw, y + 87), color, None, 12)
            d.text((xx + 16, y + 43), f"{name}   {amount} units", font=font(17, "bold"), fill=WHITE)
            xx += barw + 18
        d.text((1272, y + 46), total, font=font(21, "bold"), fill=INK)
    rounded(d, (105, 1000, 1695, 1060), "#F2F6FA", None, 16)
    d.text((133, 1019), "Results return in original input order  ·  oversized jobs receive the normal planner's refusal  ·  estimates are guardrails, not OS isolation", font=font(15, "semibold"), fill="#51667D")
    return save_image(im, "qscope-figure-04-parallel-scheduler-v3.png")


def fig_benchmark() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   LOCAL BENCHMARK")
    title(d, "Measured runtime rises with circuit size")
    rounded(d, (78, 215, 1722, 1012), "#FBFCFE", "#D9E3EC", 26, 2)
    plot = (204, 340, 1575, 864)
    x0, y0, x1, y1 = plot
    # Plot uses log-scale milliseconds so small and large observations remain legible.
    minv, maxv = 1, 300
    def ymap(v):
        return y1 - (math.log10(v) - math.log10(minv)) / (math.log10(maxv) - math.log10(minv)) * (y1 - y0)
    for val in [1, 3, 10, 30, 100, 300]:
        yy = ymap(val)
        d.line((x0, yy, x1, yy), fill="#E5EBF1", width=2)
        d.text((124, yy - 15), f"{val} ms", font=font(17), fill=MUTED)
    d.text((x0, 277), "Median wall-clock runtime · logarithmic y-axis", font=font(20, "semibold"), fill=INK)
    qubits = [6, 8, 10, 12, 14]
    runtimes = [6.69, 16.04, 16.51, 57.64, 229.93]
    spreads = [3.23, 4.02, 2.52, 4.22, 17.93]
    xs = [x0 + 120 + i * 300 for i in range(5)]
    points = [(x, ymap(v)) for x, v in zip(xs, runtimes)]
    for (xa, ya), (xb, yb) in zip(points, points[1:]):
        d.line((xa, ya, xb, yb), fill=BLUE, width=7)
    for i, (x, val, spread) in enumerate(zip(xs, runtimes, spreads)):
        yy = ymap(val)
        d.ellipse((x - 13, yy - 13, x + 13, yy + 13), fill=TEAL, outline=WHITE, width=3)
        d.text((x - 54, yy - 54 if yy - 54 > y0 else yy + 25), f"{val:.2f} ms", font=font(18, "bold"), fill=INK)
        d.text((x - 32, y1 + 22), f"{qubits[i]} qubits", font=font(18, "semibold"), fill=INK)
    d.text((x0 + 680, 912), "Generated 6-layer circuits · 512 shots · 5 measured trials · 1 warm-up", font=font(17, "semibold"), fill="#526A80")
    d.text((x0 + 680, 948), "Windows 11 · Python 3.12.14 · NumPy 2.3.5 · 16 logical CPUs · no GPU", font=font(17), fill=MUTED)
    d.text((x0 + 680, 984), "One machine / one circuit family; descriptive measurements, not a universal limit or speedup claim.", font=font(16), fill=MUTED)
    return save_image(im, "qscope-figure-05-benchmark-v3.png")


def arrow(draw, start, end, color=TEAL, width=5):
    draw.line((start[0], start[1], end[0], end[1]), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    points = [
        end,
        (end[0] - size * math.cos(angle - 0.55), end[1] - size * math.sin(angle - 0.55)),
        (end[0] - size * math.cos(angle + 0.55), end[1] - size * math.sin(angle + 0.55)),
    ]
    draw.polygon(points, fill=color)


def diagram_card(draw, bounds, heading, body, accent=TEAL, *, fill=WHITE):
    x1, y1, x2, y2 = bounds
    rounded(draw, bounds, fill, LINE, 22, 2)
    draw.rounded_rectangle((x1, y1, x2, y1 + 13), radius=7, fill=accent)
    draw.text((x1 + 24, y1 + 33), heading, font=font(22, "bold"), fill=INK)
    paragraph(draw, (x1 + 24, y1 + 76), body, font(17), "#43566B", x2 - x1 - 48, 7)


def fig_workflow() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   PROJECT WORKFLOW")
    title(d, "From a question to a reproducible experiment")
    stages = [
        ("01  START", "Choose a learning goal, a guided template, an algorithm, imported QASM, or a blank circuit.", BLUE),
        ("02  BUILD", "Compose and inspect the circuit. The engine validates gates, wires, and circuit resources.", TEAL),
        ("03  PLAN + RUN", "Check memory requirements, select a supported simulation path, and execute locally.", VIOLET),
        ("04  UNDERSTAND", "Read outcomes and probabilities, trace gate effects, and explore analysis or optimization.", AMBER),
        ("05  RETURN", "Reproduce, compare, or extend a saved run; export a report when ready.", "#D66C84"),
    ]
    x0, gap, boxw = 66, 34, 310
    for index, (heading, body, accent) in enumerate(stages):
        x = x0 + index * (boxw + gap)
        diagram_card(d, (x, 330, x + boxw, 725), heading, body, accent)
        if index < len(stages) - 1:
            arrow(d, (x + boxw + 3, 522), (x + boxw + gap - 4, 522))
    rounded(d, (100, 828, 1700, 1014), "#E9F6F3", "#B8E4D9", 20, 2)
    d.text((132, 852), "ONE SHARED EXPERIMENT CONTEXT", font=font(19, "bold"), fill="#168E78")
    paragraph(d, (132, 895), "Circuit + settings + result provenance move with the user between Build, Execute, Trace, Analyze, History, and Research.", font(19), "#43566B", 1520, 9)
    return save_image(im, "qscope-figure-01-workflow.png")


def fig_architecture_flow() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   SOFTWARE ARCHITECTURE")
    title(d, "The interface and engine meet through a local API")
    cards = [
        (70, 305, 440, 680, "FRONTEND", "React + TypeScript; circuit editor; guided modes; canvas; state and trace views; history; reports; chat.", BLUE),
        (510, 305, 850, 680, "API LAYER", "FastAPI + Pydantic; HTTP endpoints and WebSockets; request validation; structured results and errors.", TEAL),
        (920, 305, 1260, 680, "QUANTUM ENGINE", "Python + NumPy; circuit model and simulation; trace; analysis; optimization; resource planning.", VIOLET),
        (1330, 305, 1730, 680, "LOCAL DATA", "SQLite + local files; experiment history; circuit snapshots and settings; report exports.", AMBER),
    ]
    for x1, y1, x2, y2, heading, body, accent in cards:
        diagram_card(d, (x1, y1, x2, y2), heading, body, accent)
    arrow(d, (444, 490), (504, 490))
    arrow(d, (854, 490), (914, 490))
    arrow(d, (1264, 490), (1324, 490))
    arrow(d, (505, 730), (445, 730), color=BLUE)
    d.text((450, 756), "result + provenance", font=font(14, "semibold"), fill=BLUE)
    rounded(d, (170, 830, 1630, 1000), "#F2F6FA", LINE, 20, 2)
    d.text((202, 853), "OPTIONAL ASSISTANT NARRATION", font=font(18, "bold"), fill="#6552C6")
    paragraph(d, (202, 892), "The assistant receives explicit QScope evidence. A configured OpenAI-compatible endpoint may rewrite the explanation; computed values remain sourced from the local engine.", font(17), "#43566B", 1380, 7)
    return save_image(im, "qscope-figure-02-architecture.png")


def fig_simulation_tradeoffs() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   SIMULATION REPRESENTATIONS")
    title(d, "Choose a state representation that fits the question")
    # Log2 storage growth: one complex value per vector entry, or per matrix entry.
    x0, y0, x1, y1 = 190, 350, 1605, 815
    for n in range(0, 13, 2):
        xx = x0 + n / 12 * (x1 - x0)
        d.line((xx, y0, xx, y1), fill="#E6ECF2", width=2)
        d.text((xx - 7, y1 + 15), str(n), font=font(16), fill=MUTED)
    for value in range(0, 25, 4):
        yy = y1 - value / 24 * (y1 - y0)
        d.line((x0, yy, x1, yy), fill="#E6ECF2", width=2)
        d.text((116, yy - 10), str(value), font=font(15), fill=MUTED)
    d.text((785, 864), "number of qubits (n)", font=font(18, "semibold"), fill=INK)
    d.text((72, 272), "log₂ of complex storage entries · asymptotic growth", font=font(18, "semibold"), fill=INK)
    def plot_line(points, color, width=6, *, dashed=False, triangle=False):
        mapped = [(x0 + n / 12 * (x1 - x0), y1 - v / 24 * (y1 - y0)) for n, v in points]
        if dashed:
            for (xa, ya), (xb, yb) in zip(mapped, mapped[1:]):
                length = math.hypot(xb - xa, yb - ya)
                for offset in range(0, int(length), 28):
                    end_offset = min(offset + 15, length)
                    t0, t1 = offset / length, end_offset / length
                    d.line((xa + (xb - xa) * t0, ya + (yb - ya) * t0,
                            xa + (xb - xa) * t1, ya + (yb - ya) * t1), fill=color, width=width)
        else:
            d.line(mapped, fill=color, width=width)
        for x, y in mapped[::2]:
            if triangle:
                d.polygon([(x, y - 9), (x - 8, y + 7), (x + 8, y + 7)], fill=color)
            else:
                d.ellipse((x - 7, y - 7, x + 7, y + 7), fill=color)
    plot_line([(n, n) for n in range(13)], BLUE)
    plot_line([(n, 2 * n) for n in range(13)], VIOLET)
    legend = [("State vector · O(2ⁿ)", BLUE), ("Density matrix · O(4ⁿ)", VIOLET), ("Trajectory · O(2ⁿ) per sampled path", TEAL)]
    # Trajectories store pure-state vectors; total work also depends on samples.
    plot_line([(n, n) for n in range(13)], TEAL, 3, dashed=True, triangle=True)
    lx = 220
    for label, color in legend:
        d.line((lx, 954, lx + 50, 954), fill=color, width=6)
        d.text((lx + 62, 941), label, font=font(16, "semibold"), fill=INK)
        lx += 500
    d.text((92, 1044), "Conceptual memory scaling, not measured runtime. Density-matrix storage grows quadratically in state-vector dimension; trajectory cost also depends on the number of samples.", font=font(14), fill=MUTED)
    return save_image(im, "qscope-figure-03-representations.png")


def fig_execution_safety() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   RESOURCE-AWARE EXECUTION")
    title(d, "Every run is planned before the simulator allocates state")
    diagram_card(d, (92, 270, 540, 466), "RUN REQUEST", "Circuit and backend preference; shots; seed; noise settings; memory budget.", BLUE)
    diagram_card(d, (684, 270, 1132, 466), "VALIDATE + ESTIMATE", "Resolve the circuit source; estimate state and peak memory; check supported engine paths.", TEAL)
    arrow(d, (545, 368), (676, 368))
    # Decision diamond
    diamond = [(900, 570), (1120, 690), (900, 810), (680, 690)]
    d.polygon(diamond, fill="#F5F2FF", outline=VIOLET)
    d.line(diamond + [diamond[0]], fill=VIOLET, width=4)
    d.text((790, 663), "Within the\nconfigured budget?", font=font(21, "bold"), fill=INK, align="center")
    arrow(d, (908, 470), (908, 562), color=VIOLET)
    diagram_card(d, (90, 850, 680, 1030), "NO · REFUSE SAFELY", "Return the estimate and limiting constraint; do not present a simulation result.", "#D66C84")
    diagram_card(d, (1120, 850, 1710, 1030), "YES · EXECUTE LOCALLY", "Select an engine and run the circuit; attach mode, seed, shots, timing, and memory provenance.", TEAL)
    arrow(d, (680, 690), (388, 842), color="#D66C84")
    arrow(d, (1120, 690), (1415, 842), color=TEAL)
    d.text((422, 722), "No", font=font(16, "bold"), fill="#D66C84")
    d.text((1285, 722), "Yes", font=font(16, "bold"), fill=TEAL)
    return save_image(im, "qscope-figure-04-execution.png")


def fig_results_lifecycle() -> bytes:
    im, d, fg, sub = canvas(label="Q S C O P E   /   RESULTS AND REPRODUCIBILITY")
    title(d, "A run becomes a record you can inspect and revisit")
    diagram_card(d, (84, 290, 470, 610), "SIMULATION RESULT", "Counts and probabilities; backend and representation; shots, seed, and noise mode; runtime and state memory.", BLUE)
    diagram_card(d, (550, 290, 936, 610), "TRACE + ANALYZE", "Gate-by-gate state changes; entanglement and state metrics; optimization with verification; stated limits.", VIOLET)
    diagram_card(d, (1016, 290, 1402, 610), "HISTORY RECORD", "Circuit snapshot; run settings and provenance; measured result; timestamp and label.", TEAL)
    diagram_card(d, (1482, 290, 1718, 610), "NEXT", "Reproduce; compare; extend; export.", AMBER)
    arrow(d, (475, 450), (544, 450))
    arrow(d, (941, 450), (1010, 450))
    arrow(d, (1407, 450), (1476, 450))
    rounded(d, (190, 760, 1610, 988), "#F2F6FA", LINE, 22, 2)
    d.text((228, 786), "RESEARCH ASSISTANT", font=font(19, "bold"), fill="#6552C6")
    paragraph(d, (228, 834), "Attach the circuit, run, trace, or stored records needed for a question. QScope answers from those artifacts, cites available evidence, and states when the data is insufficient.", font(18), "#43566B", 1340, 8)
    arrow(d, (1220, 620), (1220, 716), color=TEAL)
    return save_image(im, "qscope-figure-05-lifecycle.png")


def insert_after(paragraph: Paragraph, *, image_bytes: bytes | None = None, caption: str | None = None) -> None:
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    inserted = Paragraph(node, paragraph._parent)
    if image_bytes is not None:
        inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
        inserted.paragraph_format.space_before = Pt(5)
        inserted.paragraph_format.space_after = Pt(2)
        inserted.paragraph_format.keep_together = True
        inserted.paragraph_format.keep_with_next = True
        inserted.add_run().add_picture(io.BytesIO(image_bytes), width=Inches(6.6))
    else:
        inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
        inserted.paragraph_format.space_before = Pt(0)
        inserted.paragraph_format.space_after = Pt(8)
        inserted.paragraph_format.keep_with_next = False
        run = inserted.add_run(caption or "")
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(74, 91, 109)


def insert_figure(doc, anchor, figure_number, caption, image_bytes):
    target = next((p for p in doc.paragraphs if anchor in p.text), None)
    if target is None:
        raise ValueError(f"Could not find insertion point: {anchor}")
    insert_after(target, image_bytes=image_bytes)
    image_p = Paragraph(target._p.getnext(), target._parent)
    insert_after(image_p, caption=f"Fig. {figure_number}. {caption}")


def editable_cell(cell, heading, body, accent, fill):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=105, start=145, bottom=105, end=145)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shade)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10" if edge == "top" else "5")
        node.set(qn("w:color"), accent)
        borders.append(node)
    cell._tc.get_or_add_tcPr().append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    run = p.add_run(heading)
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.keep_together = True
    run = p2.add_run(body)
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(INK)


def editable_arrow(cell):
    cell.text = "↓"
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)


def insert_editable_steps(doc, anchor, figure_number, caption, steps, *, note=None):
    target = next((p for p in doc.paragraphs if anchor in p.text), None)
    if target is None:
        raise ValueError(f"Could not find insertion point: {anchor}")
    table = doc.add_table(rows=len(steps) * 2 - 1 + (1 if note else 0), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(6.85)
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    for index, (heading, body, accent, fill) in enumerate(steps):
        editable_cell(table.cell(index * 2, 0), heading, body, accent, fill)
        if index < len(steps) - 1:
            editable_arrow(table.cell(index * 2 + 1, 0))
    if note:
        editable_cell(table.cell(len(steps) * 2 - 1, 0), note[0], note[1], note[2], note[3])
    table._tbl.getparent().remove(table._tbl)
    target._p.addnext(table._tbl)
    caption_node = OxmlElement("w:p")
    table._tbl.addnext(caption_node)
    caption_p = Paragraph(caption_node, target._parent)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(3)
    caption_p.paragraph_format.space_after = Pt(8)
    caption_p.paragraph_format.keep_with_next = False
    run = caption_p.add_run(f"Fig. {figure_number}. {caption}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(74, 91, 109)


def insert_editable_comparison(doc, anchor, figure_number, caption):
    target = next((p for p in doc.paragraphs if anchor in p.text), None)
    if target is None:
        raise ValueError(f"Could not find insertion point: {anchor}")
    headers = ["Representation", "State model", "Memory growth", "Useful for"]
    rows = [
        ["State vector", "Pure-state amplitude vector", "O(2ⁿ) complex values", "Ideal evolution and sampled pure-state paths"],
        ["Density matrix", "Mixed-state density operator", "O(4ⁿ) complex values", "Exact mixed-state evolution for supported noise"],
        ["Trajectory", "Sampled pure-state paths", "O(2ⁿ) per path; work grows with samples", "Approximate noisy evolution when exact matrices are costly"],
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.25, 1.85, 1.55, 2.2]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.text = text
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shade)
        set_cell_margins(cell, top=90, start=90, bottom=90, end=90)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col, text in enumerate(values):
            cell = cells[col]
            cell.width = Inches(widths[col])
            cell.text = text
            set_cell_margins(cell, top=90, start=90, bottom=90, end=90)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F3F7FB" if index % 2 == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string(INK)
    table._tbl.getparent().remove(table._tbl)
    target._p.addnext(table._tbl)
    caption_node = OxmlElement("w:p")
    table._tbl.addnext(caption_node)
    caption_p = Paragraph(caption_node, target._parent)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(3)
    caption_p.paragraph_format.space_after = Pt(8)
    run = caption_p.add_run(f"Fig. {figure_number}. {caption}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(74, 91, 109)


def clean_manuscript(doc):
    heading_map = {
        "INTRODUCTION": "I. INTRODUCTION",
        "BACKGROUND AND RELATED WORK": "II. BACKGROUND AND RELATED WORK",
        "SYSTEM DESIGN": "III. SYSTEM DESIGN",
        "IV. IMPLEMENTATION": "IV. IMPLEMENTATION",
        "EVALUATION METHODOLOGY": "V. EVALUATION METHODOLOGY",
        "RESULTS": "VI. RESULTS",
        "DISCUSSION": "VII. DISCUSSION",
        "LIMITATIONS AND FUTURE WORK": "VIII. LIMITATIONS AND FUTURE WORK",
        "CONCLUSION": "IX. CONCLUSION",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if text.startswith("### "):
            clean = text[4:]
            paragraph.style = doc.styles["Heading 2"]
            paragraph.text = clean
        elif text.startswith("* ") and paragraph.style.name == "Normal":
            paragraph.text = text[2:]
            paragraph.style = doc.styles["List Bullet"]
        elif text in heading_map and paragraph.style.name.startswith("Heading"):
            paragraph.text = heading_map[text]
        elif text in {"[K Rithvik]", "[C-Tech Department, SRM Institution of Technology, Chennai, India]"}:
            paragraph.text = text[1:-1]
        # Remove literal Markdown code ticks while leaving the actual wording intact.
        if "`" in paragraph.text:
            paragraph.text = paragraph.text.replace("`", "")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {OUTPUT}")

    figures = [
        fig_workflow(),
        fig_architecture_flow(),
        fig_simulation_tradeoffs(),
        fig_execution_safety(),
        fig_results_lifecycle(),
    ]
    doc = Document(SOURCE)
    clean_manuscript(doc)
    anchors = [
        "The central engineering goal is to make execution assumptions visible",
        "The API is hosted locally; no cloud service or quantum-hardware provider is required",
        "not by a claim of universal device-level fidelity",
        "not a hard operating-system isolation boundary",
        "Across the tested points, the median increased",
    ]
    captions = [
        "QScope user workflow from choosing a starting point to returning to saved, reproducible work.",
        "Local software architecture and data flow among the React frontend, FastAPI layer, Python/NumPy engine, and SQLite-backed storage.",
        "Conceptual asymptotic memory growth for the state-vector, density-matrix, and trajectory representations; this chart is not a runtime measurement.",
        "Resource-aware execution path: validate and estimate first, refuse requests that exceed the configured budget, and attach provenance to completed runs.",
        "Lifecycle of simulation results through trace and analysis, stored history, reproduction, comparison, export, and evidence-grounded research assistance.",
    ]
    for index, (anchor, caption) in enumerate(zip(anchors, captions), start=1):
        insert_figure(doc, anchor, index, caption, figures[index - 1])
    doc.core_properties.title = "QScope: A Local-First, Resource-Aware Quantum Circuit Simulation and Research Workspace"
    doc.core_properties.subject = "Research manuscript with explanatory workflow, architecture, simulation, execution, and reproducibility diagrams"
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Saved 5 explanatory figures to {FIGURE_DIR}")


if __name__ == "__main__":
    main()
