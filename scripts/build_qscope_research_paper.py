"""Render the QScope research-paper Markdown as a two-column IEEE-style DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "QScope_Research_Paper.md"
OUTPUT = ROOT / "docs" / "QScope_Research_Paper.docx"
FONT = "Times New Roman"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_columns(section, count: int) -> None:
    columns = section._sectPr.xpath("./w:cols")
    if columns:
        element = columns[0]
    else:
        element = OxmlElement("w:cols")
        section._sectPr.append(element)
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), "360")


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, *, size: float | None = None) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    set_run_font(run, size=8)


def add_table(document, rows: list[list[str]], table_index: int) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index in range(width):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row[c_index].strip() if c_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=7.0 if width > 4 else 7.5)
            if r_index == 0:
                set_cell_shading(cell, "DCE6F1")
                for run in paragraph.runs:
                    run.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)


def main() -> None:
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = source_lines[0].removeprefix("# ").strip()
    author_line = source_lines[2].strip().removeprefix("**").removesuffix("**")
    affiliation = source_lines[3].strip()
    email = source_lines[4].strip()

    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = "Local-first quantum circuit simulation and research workspace"
    document.core_properties.author = "[Author name — replace before submission]"
    document.core_properties.keywords = "quantum circuit simulation, reproducibility, resource-aware scheduling"

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.63)
        section.right_margin = Inches(0.63)
        set_columns(section, 1)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading1 = document.styles["Heading 1"]
    heading1.font.name = FONT
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading1.font.size = Pt(9.5)
    heading1.font.bold = True
    heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading1.paragraph_format.keep_with_next = True
    heading1.paragraph_format.space_before = Pt(7)
    heading1.paragraph_format.space_after = Pt(3)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = FONT
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading2.font.size = Pt(9)
    heading2.font.bold = True
    heading2.font.italic = True
    heading2.paragraph_format.keep_with_next = True
    heading2.paragraph_format.space_before = Pt(4)
    heading2.paragraph_format.space_after = Pt(2)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(9)
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=20, bold=True)

    for text, size, bold in ((author_line, 11, True), (affiliation, 9, False), (email, 8.5, False)):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)

    title_block_end = False
    index = 6
    body_lines = source_lines[index:]
    index = 0
    table_counter = 0
    while index < len(body_lines):
        line = body_lines[index].strip()
        if not line or line == "---":
            index += 1
            continue
        if line == "<!-- TWO_COLUMN -->":
            section = document.add_section(WD_SECTION_START.CONTINUOUS)
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)
            section.top_margin = Inches(0.68)
            section.bottom_margin = Inches(0.78)
            section.left_margin = Inches(0.63)
            section.right_margin = Inches(0.63)
            set_columns(section, 2)
            title_block_end = True
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1" if re.match(r"## [IVX]+\.", line) else "Heading 2")
            add_inline(paragraph, line[3:])
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[list[str]] = []
            while index < len(body_lines) and body_lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in body_lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row):
                    table_lines.append(row)
                index += 1
            table_counter += 1
            add_table(document, table_lines, table_counter)
            continue
        if re.match(r"\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
            index += 1
            continue
        if line.startswith("**Abstract—**") or line.startswith("**Keywords—**"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, line, size=8.5)
            index += 1
            continue
        if line.startswith("**TABLE "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, line, size=8)
            index += 1
            continue
        if line.startswith("[") and re.match(r"\[\d+\]", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, line, size=8)
            index += 1
            continue

        paragraph = document.add_paragraph()
        if title_block_end:
            paragraph.paragraph_format.first_line_indent = Inches(0.14)
        add_inline(paragraph, line)
        index += 1

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.space_after = Pt(0)
        header_run = header.add_run("QScope  |  Research Paper")
        set_run_font(header_run, size=8, italic=True)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(footer)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
